from pathlib import Path
from pypdf import PdfReader
from docx import Document
import openpyxl


class FileLoaderService:
    def load_file(self, path: str) -> dict:
        p = Path(path)
        suffix = p.suffix.lower()
        if suffix == ".pdf":
            return {"title": p.name, "content_markdown": self._pdf(p)}
        if suffix == ".docx":
            return {"title": p.name, "content_markdown": self._docx(p)}
        if suffix in [".xlsx", ".xlsm"]:
            return {"title": p.name, "content_markdown": self._xlsx(p)}
        return {"title": p.name, "content_markdown": p.read_text(encoding="utf-8", errors="ignore")}

    def _pdf(self, p: Path) -> str:
        reader = PdfReader(str(p))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    def _docx(self, p: Path) -> str:
        doc = Document(str(p))
        return "\n".join(x.text for x in doc.paragraphs if x.text.strip())

    def _xlsx(self, p: Path) -> str:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines)

file_loader_service = FileLoaderService()
